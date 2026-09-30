"""
Column display names, managed on the admin screen.

Which Korean name a column shows is data, not code: an override per connection
(and per table where a name must differ), falling back to the DB comment. The
computed-column terms (합계, 건수 ...) are managed here too.
"""
from io import BytesIO
from typing import List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.deps import ADMIN
from app.db.column_labels import (
    DEFAULT_EXPRESSION_TERMS,
    Override,
    apply_labels,
    coverage,
    merge_expression_terms,
    parse_label_sheet,
)
from app.db.connection_pool import DatabaseConnectionPool
from app.db.repositories.column_labels import ColumnLabelRepository
from app.db.repositories.database_repo import DatabaseRepository
from app.dependencies import get_db, get_db_pool

router = APIRouter(tags=["Column labels"])

MAX_SHEET_BYTES = 5 * 1024 * 1024
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class LabelIn(BaseModel):
    column_name: str = Field(..., min_length=1, max_length=128)
    # None: the column name on the whole connection. A value: that table only.
    table_key: Optional[str] = Field(None, max_length=255)
    # Blank removes the override, so the DB comment shows again.
    label: str = Field("", max_length=100)


class TermIn(BaseModel):
    label: str = Field(..., min_length=1, max_length=50)


async def _connection_schema(connection_id: int, db: AsyncSession, pool: DatabaseConnectionPool):
    conn = await DatabaseRepository(db).get_by_id(connection_id)
    if conn is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Connection {connection_id} not found")
    try:
        return await pool.get_schema(str(connection_id))
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"스키마를 읽지 못했습니다: {exc}")


def _problem_with(schema, table_key: Optional[str], column_name: str) -> Optional[str]:
    """Why this label cannot be attached, or None. A label for a column the
    connection does not have would silently never show."""
    if table_key is not None:
        if table_key not in schema:
            return f"'{table_key}' 테이블이 이 연결에 없습니다"
        if column_name not in {c["name"] for c in schema[table_key]}:
            return f"'{table_key}'에 '{column_name}' 컬럼이 없습니다"
    elif not any(c["name"] == column_name for cols in schema.values() for c in cols):
        return f"'{column_name}' 컬럼이 이 연결의 어떤 테이블에도 없습니다"
    return None


@router.get("/databases/{connection_id}/column-labels", dependencies=[ADMIN])
async def list_labels(
    connection_id: int,
    db: AsyncSession = Depends(get_db),
    pool: DatabaseConnectionPool = Depends(get_db_pool),
):
    """Every distinct column name with its effective label, plus how many have none."""
    schema = await _connection_schema(connection_id, db, pool)
    columns = coverage(schema, await ColumnLabelRepository(db).overrides(connection_id))
    mapped = sum(1 for c in columns if c["mapped"])
    return {
        "columns": columns,
        "tables": sorted(schema.keys()),
        "summary": {"total": len(columns), "mapped": mapped, "unmapped": len(columns) - mapped},
    }


@router.put("/databases/{connection_id}/column-labels", dependencies=[ADMIN])
async def set_label(
    connection_id: int,
    body: LabelIn,
    db: AsyncSession = Depends(get_db),
    pool: DatabaseConnectionPool = Depends(get_db_pool),
):
    """Set a column's name, or clear it with a blank label."""
    schema = await _connection_schema(connection_id, db, pool)
    repo = ColumnLabelRepository(db)
    label = body.label.strip()

    if not label:
        await repo.clear(connection_id, body.table_key, body.column_name)
        return {"cleared": True}

    problem = _problem_with(schema, body.table_key, body.column_name)
    if problem:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=problem)
    await repo.set(connection_id, body.table_key, body.column_name, label)
    return {"cleared": False, "label": label}


@router.delete("/databases/{connection_id}/column-labels/{label_id}", status_code=status.HTTP_204_NO_CONTENT, dependencies=[ADMIN])
async def delete_label(connection_id: int, label_id: int, db: AsyncSession = Depends(get_db)):
    if not await ColumnLabelRepository(db).delete(connection_id, label_id):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="한글명을 찾을 수 없습니다")


@router.get("/databases/{connection_id}/column-labels/export", dependencies=[ADMIN])
async def export_labels(
    connection_id: int,
    db: AsyncSession = Depends(get_db),
    pool: DatabaseConnectionPool = Depends(get_db_pool),
):
    """The mapping as a workbook, unmapped columns included with a blank name to
    fill in. It is the same layout the import reads."""
    schema = await _connection_schema(connection_id, db, pool)
    rows = coverage(schema, await ColumnLabelRepository(db).overrides(connection_id))

    book = Workbook()
    sheet = book.active
    sheet.title = "컬럼 한글명"
    sheet.append(["컬럼명", "한글명", "테이블(비우면 전체)", "DB 설명", "쓰이는 테이블"])
    for row in sorted(rows, key=lambda r: (r["mapped"], r["name"])):
        sheet.append([row["name"], row["label"] or "", "", row["comment"] or "", ", ".join(row["tables"])])
        for exc in row["exceptions"]:
            sheet.append([row["name"], exc["label"], exc["table_key"], row["comment"] or "", exc["table_key"]])
    for letter, width in zip("ABCDE", (24, 28, 24, 40, 40)):
        sheet.column_dimensions[letter].width = width

    out = BytesIO()
    book.save(out)
    return Response(
        content=out.getvalue(),
        media_type=XLSX,
        headers={"Content-Disposition": f"attachment; filename=column_labels_{connection_id}.xlsx"},
    )


@router.post("/databases/{connection_id}/column-labels/import", dependencies=[ADMIN])
async def import_labels(
    connection_id: int,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    pool: DatabaseConnectionPool = Depends(get_db_pool),
):
    """Apply a sheet of (컬럼명, 한글명[, 테이블]) rows.

    Rows that would change nothing (the name already shows) are not stored, so
    importing an export does not freeze every DB comment as an override. Rows the
    connection cannot use are reported and skipped; the rest are applied.
    """
    schema = await _connection_schema(connection_id, db, pool)
    raw = await file.read()
    if len(raw) > MAX_SHEET_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="파일이 5MB를 넘습니다")
    try:
        book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        rows = [list(r) for r in book.worksheets[0].iter_rows(values_only=True)]
    except Exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="엑셀(.xlsx) 파일을 읽지 못했습니다")

    items, problems = parse_label_sheet(rows)
    repo = ColumnLabelRepository(db)
    current = apply_labels(schema, await repo.overrides(connection_id))
    shown_now = {(t, c["name"]): c["label"] for t, cols in current.items() for c in cols}

    to_apply: List[Override] = []
    unchanged = 0
    for item in items:
        reason = _problem_with(schema, item.table_key, item.column_name)
        if reason:
            problems.append({"row": None, "column": item.column_name, "reason": reason})
            continue
        tables = [item.table_key] if item.table_key else [t for t, cols in schema.items()
                                                          if any(c["name"] == item.column_name for c in cols)]
        if all(shown_now.get((t, item.column_name)) == item.label for t in tables):
            unchanged += 1
            continue
        to_apply.append(item)

    applied = await repo.set_many(connection_id, to_apply)
    return {"applied": applied, "unchanged": unchanged, "problems": problems}


# --- computed-column terms (global) ---------------------------------------------------------

@router.get("/expression-terms")
async def list_terms(db: AsyncSession = Depends(get_db)):
    """Names for SUM/COUNT/... columns, defaults merged with any change."""
    return merge_expression_terms(await ColumnLabelRepository(db).terms())


@router.put("/expression-terms/{func}", dependencies=[ADMIN])
async def set_term(func: str, body: TermIn, db: AsyncSession = Depends(get_db)):
    if func not in DEFAULT_EXPRESSION_TERMS:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"'{func}'는 없는 항목입니다")
    label = body.label.strip()
    if not label:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="이름을 입력하세요")
    repo = ColumnLabelRepository(db)
    if label == DEFAULT_EXPRESSION_TERMS[func]:
        await repo.reset_term(func)  # same as the default: nothing to keep
    else:
        await repo.set_term(func, label)
    return merge_expression_terms(await repo.terms())


@router.delete("/expression-terms/{func}", dependencies=[ADMIN])
async def reset_term(func: str, db: AsyncSession = Depends(get_db)):
    if func not in DEFAULT_EXPRESSION_TERMS:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"'{func}'는 없는 항목입니다")
    repo = ColumnLabelRepository(db)
    await repo.reset_term(func)
    return merge_expression_terms(await repo.terms())
